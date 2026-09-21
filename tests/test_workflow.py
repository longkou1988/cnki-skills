import importlib.util
import json
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'skills/cnki-resume/scripts/workflow.py'
spec = importlib.util.spec_from_file_location('workflow', SCRIPT)
w = importlib.util.module_from_spec(spec)
spec.loader.exec_module(w)
spec2 = importlib.util.spec_from_file_location('installer2', ROOT / 'scripts/install.py')
i = importlib.util.module_from_spec(spec2)
spec2.loader.exec_module(i)


def rec(**changes):
    value = dict(title='合成论文：数字化与创新', url='https://example.org/synthetic',
                 database='SYNTHETIC', record_id='1', retrieved_at='2026-09-21', year='2025')
    value.update(changes)
    return value


def decision(kind='include'):
    return dict(decision=kind, reason='符合样本标准（合成）', evidence='中国企业样本（合成）',
                basis='abstract', access_level='abstract_only')


def extraction():
    return dict(access_level='abstract_only', fields={'sample': dict(value='中国企业（合成）',
        quote='中国企业样本（合成）', locator='摘要', source='https://example.org/synthetic'), 'method': None})


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.path = self.root / 'task.sqlite'
        w.init(self.path, dict(topic='合成测试', query='合成', criteria='中国企业', stages=list(w.STAGES)))
        self.db = w.connect(self.path)
        with self.db:
            self.pid = w.import_records(self.db, [rec()])[0]['id']

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def test_no_overwrite_or_create_on_resume(self):
        with self.assertRaises(FileExistsError):
            w.init(self.path, dict(topic='x', query='x', criteria='x'))
        absent = self.root / 'absent.sqlite'
        with self.assertRaises(ValueError): w.connect(absent)
        self.assertFalse(absent.exists())

    def test_duplicate_and_atomic_conflict(self):
        with self.db:
            r = w.import_records(self.db, [rec()])
        self.assertEqual(r[0]['status'], 'existing')
        with self.assertRaises(ValueError), self.db:
            w.import_records(self.db, [rec(record_id='2'), rec(title='冲突标题')])
        self.assertEqual(w.status(self.db)['papers'], 1)

    def test_distinct_ids_and_doi_normalization(self):
        with self.db:
            result = w.import_records(self.db, [rec(record_id='2'), rec(record_id='3', doi='https://doi.org/10.1234/ABC'), rec(record_id='3', doi='10.1234/abc')])
        self.assertEqual(result[0]['possible_duplicates'], [self.pid])
        self.assertEqual(result[2]['status'], 'existing')
        self.assertEqual(w.status(self.db)['papers'], 3)

    def test_enrichment_preserves_identity_and_completed_work(self):
        with self.db:
            w.screen(self.db, self.pid, decision())
            r = w.import_records(self.db, [rec(doi='10.1234/NEW', abstract='新增摘要')])
            r2 = w.import_records(self.db, [rec(doi='10.1234/new', record_id='', database='')])
        self.assertEqual(r[0]['id'], self.pid)
        self.assertEqual(r2[0]['id'], self.pid)
        self.assertEqual(w.status(self.db)['papers'], 1)
        self.assertEqual(w.stage(self.db, self.pid, 'screening')['status'], 'done')
        self.assertEqual(json.loads(w.paper(self.db, self.pid)['metadata'])['abstract'], '新增摘要')
        with self.assertRaises(ValueError), self.db:
            w.import_records(self.db, [rec(doi='10.1234/different')])

    def test_rollback_invalid_evidence(self):
        with self.db: w.screen(self.db, self.pid, decision())
        e = extraction(); del e['fields']['sample']['quote']
        with self.assertRaises(ValueError), self.db: w.extract(self.db, self.pid, e)
        self.assertIsNone(w.paper(self.db, self.pid)['extraction'])
        self.assertEqual(w.stage(self.db, self.pid, 'extraction')['status'], 'pending')

    def test_access_level_and_unknown_fields(self):
        d = decision(); d['basis'] = 'full_text'
        with self.assertRaises(ValueError), self.db: w.screen(self.db, self.pid, d)
        with self.db: w.screen(self.db, self.pid, decision())
        for field in ('unknown', 'sample'):
            e = extraction()
            if field == 'unknown': e['fields'][field] = None
            else: e['fields']['sample']['locator'] = 'PDF第3页'
            with self.assertRaises(ValueError), self.db: w.extract(self.db, self.pid, e)

    def test_decision_change_invalidates_extraction_but_repeat_does_not(self):
        with self.db:
            w.screen(self.db, self.pid, decision())
            w.extract(self.db, self.pid, extraction())
            w.screen(self.db, self.pid, decision())
        self.assertEqual(w.stage(self.db, self.pid, 'extraction')['status'], 'done')
        with self.db: w.screen(self.db, self.pid, decision('exclude'))
        self.assertIsNone(w.paper(self.db, self.pid)['extraction'])
        self.assertEqual(w.stage(self.db, self.pid, 'extraction')['status'], 'skipped')
        self.assertTrue(self.db.execute("SELECT detail FROM events WHERE action='extraction-evidence'").fetchone())
        with self.db: w.screen(self.db, self.pid, decision('uncertain'))
        self.assertEqual(w.stage(self.db, self.pid, 'extraction')['status'], 'blocked')

    def test_resume_does_not_repeat_done_or_running(self):
        with self.db:
            w.screen(self.db, self.pid, decision())
            w.mark(self.db, self.pid, 'download', 'running')
        self.db.close(); self.db = w.connect(self.path)
        with self.db: queue = w.queue(self.db)
        self.assertNotIn('screening', [x['stage'] for x in queue])
        self.assertEqual(next(x for x in queue if x['stage'] == 'download')['next_action'], 'inspect_active_operation')
        with self.assertRaises(ValueError), self.db: w.mark(self.db, self.pid, 'download', 'running')
        with self.db:
            w.mark(self.db, self.pid, 'download', 'pending', reason='Confirmed transfer stopped')
            w.mark(self.db, self.pid, 'download', 'running')

    def test_concurrent_claim_only_one_succeeds(self):
        command = [sys.executable, str(SCRIPT), '--task', str(self.path), 'mark',
                   '--id', self.pid, '--stage', 'download', '--status', 'running']
        processes = [subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE) for _ in range(2)]
        for process in processes: process.communicate(timeout=10)
        self.assertEqual(sorted(p.returncode for p in processes), [0, 1])

    def test_required_full_text_keeps_partial_extraction_blocked(self):
        task = self.root / 'full.sqlite'
        w.init(task, dict(topic='x', query='x', criteria='x', required_access='full_text'))
        db = w.connect(task)
        try:
            with db:
                pid = w.import_records(db, [rec()])[0]['id']
                w.screen(db, pid, decision())
                w.extract(db, pid, extraction())
            self.assertEqual(w.stage(db, pid, 'extraction')['status'], 'blocked')
            self.assertIsNotNone(w.paper(db, pid)['extraction'])
        finally:
            db.close()

    def test_no_evidence_free_completion_or_premature_extraction(self):
        with self.assertRaises(ValueError), self.db: w.mark(self.db, self.pid, 'screening', 'done')
        with self.assertRaises(ValueError), self.db: w.mark(self.db, self.pid, 'extraction', 'running')
        with self.assertRaises(ValueError), self.db: w.extract(self.db, self.pid, extraction())

    def test_download_validation_and_hash_recheck(self):
        path = self.root / 'fake.pdf'; path.write_text('<html>login</html>')
        with self.assertRaises(ValueError), self.db: w.mark(self.db, self.pid, 'download', 'done', path=path, identity_verified=True)
        path.write_bytes(b'%PDF-1.7\nsynthetic header fixture only')
        with self.assertRaises(ValueError), self.db: w.mark(self.db, self.pid, 'download', 'done', path=path)
        with self.db: w.mark(self.db, self.pid, 'download', 'done', path=path, identity_verified=True)
        with self.db: self.assertNotIn('download', [x['stage'] for x in w.queue(self.db)])
        path.write_bytes(b'%PDF-1.7\nchanged')
        with self.db: self.assertIn('download', [x['stage'] for x in w.queue(self.db)])
        self.assertEqual(w.stage(self.db, self.pid, 'download')['status'], 'pending')

    def test_missing_file_requeued(self):
        path = self.root / 'refs.bib'; path.write_text('synthetic citation')
        with self.db: w.mark(self.db, self.pid, 'citation', 'done', path=path, identity_verified=True)
        path.unlink()
        with self.db: w.queue(self.db)
        self.assertEqual(w.stage(self.db, self.pid, 'citation')['status'], 'pending')

    def test_excel_independent_reader_unknowns_and_formula_safety(self):
        with self.db:
            w.screen(self.db, self.pid, decision())
            e = extraction(); e['fields']['sample']['value'] = '=HYPERLINK("https://example.org")'
            w.extract(self.db, self.pid, e)
            w.export(self.db, self.root / 'out')
        wb = load_workbook(self.root / 'out/研究工作表.xlsx')
        self.assertEqual(len(wb.sheetnames), 5)
        ws = wb['实证研究对照表']
        values = list(ws.values)
        self.assertIn('未提取/证据不足', values[1])
        cell = next(c for c in ws[2] if c.value == e['fields']['sample']['value'])
        self.assertEqual(cell.data_type, 's')
        self.assertEqual(ws.freeze_panes, 'A2')
        self.assertEqual(wb['原文证据']['F2'].value, '摘要')
        self.assertIn("'=HYPERLINK", (self.root / 'out/实证研究对照表.csv').read_text(encoding='utf-8-sig'))
        with self.assertRaises(FileExistsError), self.db: w.export(self.db, self.root / 'out')

    def test_cli_persistence_and_bad_input(self):
        base = [sys.executable, str(SCRIPT), '--task', str(self.path)]
        p = subprocess.run(base + ['queue'], capture_output=True, text=True)
        self.assertEqual(p.returncode, 0, p.stderr)
        self.assertEqual(len(json.loads(p.stdout)), 4)
        bad = self.root / 'bad.json'; bad.write_text('not JSON')
        p = subprocess.run(base + ['import', '--input', str(bad)], capture_output=True, text=True)
        self.assertNotEqual(p.returncode, 0)
        self.assertEqual(w.status(self.db)['papers'], 1)

    def test_credential_url_rejected(self):
        for url in ('javascript:foo', 'https://u:p@example.org', 'https://example.org/?token=abc'):
            with self.assertRaises(ValueError), self.db: w.import_records(self.db, [rec(url=url)])

    def test_invalid_xml_does_not_leave_report(self):
        with self.db:
            w.screen(self.db, self.pid, decision())
            e = extraction(); e['fields']['sample']['value'] = '\x00'
            w.extract(self.db, self.pid, e)
        with self.assertRaises(ValueError), self.db: w.export(self.db, self.root / 'bad-out')
        self.assertFalse((self.root / 'bad-out').exists())


class UpgradeTests(unittest.TestCase):
    def test_upgrade_backs_up_existing_and_keeps_unrelated(self):
        with tempfile.TemporaryDirectory() as tmp:
            i.install('codex', tmp)
            base = Path(tmp) / '.agents'
            skill = base / 'skills/cnki-screening/SKILL.md'
            skill.write_text('local customization')
            other = base / 'skills/unrelated'; other.mkdir()
            (other / 'keep').write_text('keep')
            i.install('codex', tmp, dry_run=True, upgrade=True)
            self.assertEqual(skill.read_text(), 'local customization')
            self.assertFalse((base / 'cnki-backups').exists())
            i.install('codex', tmp, upgrade=True)
            backups = list((base / 'cnki-backups').glob('*/skills/cnki-screening/SKILL.md'))
            self.assertEqual(len(backups), 1)
            self.assertEqual(backups[0].read_text(), 'local customization')
            self.assertEqual((other / 'keep').read_text(), 'keep')
            self.assertIn('cnki-screening', skill.read_text())

    def test_upgrade_refuses_symlinks(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / '.agents/skills/cnki-screening'
            target.parent.mkdir(parents=True)
            target.symlink_to(Path(tmp), target_is_directory=True)
            with self.assertRaises(FileExistsError): i.install('codex', tmp, upgrade=True)


if __name__ == '__main__': unittest.main()
